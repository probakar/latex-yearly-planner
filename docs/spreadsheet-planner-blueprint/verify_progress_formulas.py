"""Validates the RECOMMENDED progress / habit / numeric-bar / streak formulas (same method as verify_formulas.py)."""
import openpyxl, formulas, sys
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "D"
# --- Task list: A=task text, B=checkbox (TRUE/FALSE). Cases in separate column blocks.
cases = {
  "empty":    [("",False)]*5,
  "none":     [("a",False),("b",False),("c",False),("",False),("",False)],
  "partial":  [("a",True),("b",False),("c",True),("",False),("",False)],
  "all":      [("a",True),("b",True),("",False),("",False),("",False)],
  "stray":    [("a",True),("",True),("",True),("",False),("",False)],   # ticks on blank rows must NOT count
}
col = 1
for name, rows in cases.items():
    for i,(t,c) in enumerate(rows,1):
        ws.cell(i,col, t if t else None); ws.cell(i,col+1,c)
    A=openpyxl.utils.get_column_letter(col); B=openpyxl.utils.get_column_letter(col+1)
    ws.cell(7,col,f'=IFERROR(SUMPRODUCT(({A}1:{A}5<>"")*({B}1:{B}5=TRUE))/SUMPRODUCT(({A}1:{A}5<>"")*1),0)')
    col += 3
# --- Habits: 4 habit rows x 7 days. row3 blank label with ticks (must be ignored)
H = [("Read",[1,1,1,0,0,0,0]),("Walk",[1,1,1,1,1,1,1]),("",[1,1,0,0,0,0,0]),("Stretch",[0,0,0,0,0,0,0])]
for r,(n,days) in enumerate(H, 12):
    ws.cell(r,1,n if n else None)
    for d,v in enumerate(days): ws.cell(r,2+d,bool(v))
    ws.cell(r,9,f'=IF(A{r}="","",COUNTIF(B{r}:H{r},TRUE)/7)')
ws["A17"]='=IFERROR(SUMPRODUCT((A12:A15<>"")*(B12:H15=TRUE))/(COUNTA(A12:A15)*7),0)'
# daily habit % per column (Mon): ticks among labelled habits / labelled habits
ws["B17"]='=IFERROR(SUMPRODUCT((A12:A15<>"")*(B12:B15=TRUE))/COUNTA(A12:A15),0)'
# --- Numeric tracker: value, min, max -> bar fraction & REPT bar; averages ignore blanks
for i,v in enumerate([6,None,10,0,-3,12],20):
    ws.cell(i,1,v)
    ws.cell(i,2,f'=IF(A{i}="","",MIN(1,MAX(0,(A{i}-$E$20)/($F$20-$E$20))))')
    ws.cell(i,3,f'=IF(B{i}="","",REPT("█",ROUND(B{i}*10,0))&REPT("░",10-ROUND(B{i}*10,0)))')
ws["E20"]=0; ws["F20"]=10
ws["A27"]='=IFERROR(AVERAGE(A20:A25),0)'      # blanks ignored
ws["A28"]='=SUM(A20:A25)'
ws["A29"]='=COUNT(A20:A25)'
# empty tracker average must be 0 not #DIV/0!
ws["A31"]='=IFERROR(AVERAGE(A40:A45),0)'
# --- Streak (consecutive TRUE counting back from last day with a tick); row of 14 days
streak = [0,1,1,0,1,1,1,1,1,0,1,1,1,0]   # ends with 0 -> current streak 0; best streak 5
for d,v in enumerate(streak): ws.cell(34,1+d,bool(v))
# RECOMMENDED streak: helper row, no array formulas.  A37: =IF(A34=TRUE,1,0)   B37: =IF(B34=TRUE,A37+1,0) ...
ws["A37"]='=IF(A34=TRUE,1,0)'
for d in range(1,14):
    c=openpyxl.utils.get_column_letter(1+d); p_=openpyxl.utils.get_column_letter(d)
    ws[f"{c}37"]=f'=IF({c}34=TRUE,{p_}37+1,0)'
ws["A38"]="=MAX(A37:N37)"      # best streak
ws["A39"]="=N37"                # current streak (value in the last day column)
wb.save("/tmp/p.xlsx")
sol = formulas.ExcelModel().loads("/tmp/p.xlsx").finish().calculate()
g = lambda a: sol[f"'[p.xlsx]D'!{a}"].value[0,0]
exp = {"A7":0,"D7":0,"G7":2/3,"J7":1.0,"M7":1.0,
       "I12":3/7,"I13":1.0,"I14":"","I15":0.0,
       "A17":(3+7+0)/(3*7),"B17":2/3,
       "B20":0.6,"B21":"","B22":1.0,"B23":0.0,"B24":0.0,"B25":1.0,
       "C20":"██████░░░░","A27":(6+10+0-3+12)/5,"A28":25,"A29":5,"A31":0}
bad=0
for k,v in exp.items():
    got=g(k)
    ok = (got==v) if isinstance(v,str) else abs(float(got)-v)<1e-9
    print(("ok  " if ok else "FAIL"),k,repr(got),"expected",repr(v)); bad+= (not ok)
for k,v in {"A38":5,"A39":0,"M37":3}.items():
    got=g(k); ok=abs(float(got)-v)<1e-9; print(("ok  " if ok else "FAIL"),k,got,"expected",v); bad+=(not ok)
print("RESULT:", "ALL PASS" if not bad else f"{bad} FAILURES"); sys.exit(1 if bad else 0)
