"""
Generates Planner.xlsx from BLUEPRINT.md (original design, not a copy of the reference).

    python3 -m venv v && ./v/bin/pip install openpyxl
    ./v/bin/python build_planner.py            # -> Planner.xlsx

Cell addresses match BLUEPRINT.md section 18 (two small deviations are listed in README_PLANNER.md).
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.properties import PageSetupProperties

OUT = sys.argv[1] if len(sys.argv) > 1 else "Planner.xlsx"

# ---------- design system (BLUEPRINT section 19) ----------
PRIMARY, MIST, ACCENT = "24566B", "E8F1F2", "F2A541"
GREEN, TRACK, DANGER = "3BA776", "E3E9EC", "D1495B"
SURFACE, BORDER, TEXT, MUTED, DIM = "F7F9FA", "D5DDE0", "1E2A30", "6B7A83", "B8C2C7"
PRIO_COL = ["F4B6C2", "F8D0A8", "FBEBA5", "CDE8B5", "BFD9F2", "D9CCEF", "DADFE2"]
CAT_COL = ["BFE3DA", "F9D5B8", "D3DCF2", "F6C9D6", "E4EDB9", "CFE6F3", "F3DFA8", "DCCFEA",
           "C6E6C9", "F5CFC3", "C9D8E4", "E9D9C4"]
FONT = "Calibri"

def fill(c): return PatternFill("solid", start_color=c, end_color=c)
thin = Side(style="thin", color=BORDER)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
UNLOCK = Protection(locked=False)

def f(size=10, bold=False, color=TEXT, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)

def style(c, *, font=None, bg=None, border=None, align=None, fmt=None, unlock=False):
    if font: c.font = font
    if bg: c.fill = fill(bg)
    if border: c.border = border
    if align: c.alignment = align
    if fmt: c.number_format = fmt
    if unlock: c.protection = UNLOCK

def inp(c, fmt=None, **kw):      # input cell: white, boxed, unlocked
    style(c, font=f(), bg="FFFFFF", border=BOX, fmt=fmt, unlock=True, **kw)

def calc(c, fmt=None, **kw):     # formula cell: grey, locked
    style(c, font=f(), bg=SURFACE, border=BOX, fmt=fmt, **kw)

def banner(ws, row, c1, c2, text_or_formula):
    for col in range(c1, c2 + 1):
        style(ws.cell(row, col), bg=PRIMARY, font=f(10, True, "FFFFFF"))
    ws.cell(row, c1).value = text_or_formula

def lab(key):                    # label override (Settings!B94:C103)
    return f'=IF(Settings!$C${94+key}="",Settings!$B${94+key},Settings!$C${94+key})'

LABELS = ["Top goals", "Monthly to-do", "Important dates", "Notes", "Schedule", "Gratitude",
          "Habits", "Numeric trackers", "Weekly priorities", "Secondary tasks"]
(L_GOALS, L_TODO, L_DATES, L_NOTES, L_SCHED, L_GRAT, L_HAB, L_NUM, L_WPRI, L_WSEC) = range(10)

MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August",
          "September", "October", "November", "December"]
DAYS = ["Sunday", "Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday"]

wb = Workbook()
ws_start = wb.active; ws_start.title = "Start"
ws_set = wb.create_sheet("Settings")
ws_mon = wb.create_sheet("Month")
ws_w = [wb.create_sheet(f"W{i}") for i in range(1, 7)]
ws_calc = wb.create_sheet("Calc")
ws_lists = wb.create_sheet("Lists")

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)

# =====================================================================
# Lists (hidden)
# =====================================================================
for i, m in enumerate(MONTHS, 1):
    ws_lists.cell(i, 1, m)
for i, d in enumerate(DAYS, 1):
    ws_lists.cell(i, 3, d)

# =====================================================================
# Settings
# =====================================================================
S = ws_set
S.column_dimensions["A"].width = 2
S.column_dimensions["B"].width = 30
for c in "CDEF":
    S.column_dimensions[c].width = 18
S["B1"] = "Settings"; S["B1"].font = f(20, True, PRIMARY)
banner(S, 2, 2, 6, "Planner setup  (white cells are yours; grey cells are automatic)")

rows = [(3, "Month", "October"), (4, "Year", 2026), (5, "Week starts on", "Monday"),
        (6, "Schedule start time", 0.25), (7, "Schedule slot (minutes)", 30),
        (8, "Progress bar colour (hex)", "#3BA776"), (9, "Progress track colour (hex)", "#E3E9EC")]
for r, label, val in rows:
    S.cell(r, 2, label).font = f(10, True)
    inp(S.cell(r, 3, val))
S["C6"].number_format = "h:mm AM/PM"
S["D6"] = "type a time, e.g. 6:00 AM"; S["D6"].font = f(9, color=MUTED, italic=True)
S["D8"] = "bars use this colour in Google Sheets (run the setup script)"; S["D8"].font = f(9, color=MUTED, italic=True)
S["D8"].alignment = Alignment(wrap_text=False)
style(S["D3"], font=f(9, color=MUTED, italic=True)); S["D3"] = "pick from the list"
S["D4"] = "1901 to 2100"; S["D4"].font = f(9, color=MUTED, italic=True)

dv = DataValidation(type="list", formula1="=Lists!$A$1:$A$12", allow_blank=False, showErrorMessage=True,
                    errorTitle="Month", error="Choose a month from the list.")
S.add_data_validation(dv); dv.add("C3")
dv = DataValidation(type="whole", operator="between", formula1="1901", formula2="2100", allow_blank=False,
                    showErrorMessage=True, errorTitle="Year", error="Enter a year from 1901 to 2100.")
S.add_data_validation(dv); dv.add("C4")
dv = DataValidation(type="list", formula1='"Sunday,Monday"', allow_blank=False, showErrorMessage=True)
S.add_data_validation(dv); dv.add("C5")
dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="0.999988", allow_blank=False,
                    showErrorMessage=True, errorTitle="Start time", error="Enter a time of day, e.g. 6:00 AM.")
S.add_data_validation(dv); dv.add("C6")
dv = DataValidation(type="list", formula1='"15,30,60"', allow_blank=False, showErrorMessage=True)
S.add_data_validation(dv); dv.add("C7")
dv = DataValidation(type="custom", formula1='=AND(LEN(C8)=7,LEFT(C8,1)="#")', allow_blank=False,
                    showErrorMessage=True, errorTitle="Colour", error="Use a hex colour like #3BA776.")
S.add_data_validation(dv); dv.add("C8"); dv.add("C9")

def list_block(top, title, n, defaults, col=2):
    banner(S, top - 1, 2, 6, title)
    for i in range(n):
        v = defaults[i] if i < len(defaults) else None
        c = S.cell(top + i, col, v); inp(c)
list_block(12, "Priorities (7 slots – colours belong to the slot)", 7,
           ["Urgent", "High", "Normal", "Low", "Waiting", "Someday", "Delegated"])
list_block(22, "Categories (12 slots – fill top-down, no gaps)", 12,
           ["Work", "Home", "Health", "Learning", "Finance", "Social", "Errands", "Rest"])
list_block(37, "Habits (21 slots – shown on every weekly tab)", 21,
           ["Drink water", "Move / exercise", "Read", "Plan tomorrow", "Sleep on time"])

banner(S, 60, 2, 6, "Numeric trackers (5)")
for j, h in enumerate(["Label", "Unit", "Min", "Max / target"]):
    c = S.cell(60, 2 + j); c.value = h
trk = [("Water", "glasses", 0, 8), ("Mood", "1-5", 1, 5), ("Sleep", "hours", 0, 9),
       ("Pages read", "pages", 0, 50), ("Exercise", "minutes", 0, 60)]
for i, t in enumerate(trk):
    for j, v in enumerate(t):
        inp(S.cell(61 + i, 2 + j, v))

banner(S, 67, 2, 6, "Weekday translations  (leave blank to use English)")
S.cell(68, 2, "English (fixed)").font = f(9, True, MUTED); S.cell(68, 3, "Your word").font = f(9, True, MUTED)
for i, d in enumerate(DAYS):
    calc(S.cell(69 + i, 2, d)); inp(S.cell(69 + i, 3))
banner(S, 77, 2, 6, "Month translations  (display only; leave blank for English)")
S.cell(78, 2, "English (fixed)").font = f(9, True, MUTED); S.cell(78, 3, "Your word").font = f(9, True, MUTED)
for i, m in enumerate(MONTHS):
    calc(S.cell(79 + i, 2, m)); inp(S.cell(79 + i, 3))
banner(S, 92, 2, 6, "Section labels  (type a replacement to rename a heading everywhere)")
S.cell(93, 2, "Default (fixed)").font = f(9, True, MUTED); S.cell(93, 3, "Your label").font = f(9, True, MUTED)
for i, t in enumerate(LABELS):
    calc(S.cell(94 + i, 2, t)); inp(S.cell(94 + i, 3))

# =====================================================================
# Calc (hidden): date engine, week ranges
# =====================================================================
C = ws_calc
C["A1"] = "Date engine – do not edit"; C["A1"].font = f(10, True, DANGER)
eng = {
    "B1": "=MATCH(Settings!$C$3,Lists!$A$1:$A$12,0)",
    "B2": "=DATE(Settings!$C$4,B1,1)",
    "B3": "=EOMONTH(B2,0)",
    "B4": "=DAY(B3)",
    "B5": '=IF(Settings!$C$5="Monday",2,1)',
    "B6": "=WEEKDAY(B2,B5)-1",
    "B7": "=B2-B6",
    "B8": "=ROUNDUP((B6+B4)/7,0)",
}
names = {"B1": "SelMonth", "B2": "FirstDay", "B3": "LastDay", "B4": "DaysInMonth", "B5": "WSType",
         "B6": "Offset", "B7": "GridStart", "B8": "WeeksNeeded"}
for a, fm in eng.items():
    C[a] = fm
    C["A" + a[1:]] = names[a]
    C[a].number_format = "d mmm yyyy" if a in ("B2", "B3", "B7") else "0"
# weekday labels for the 7 grid columns (translation-aware): row 1, D..J
for k in range(7):
    col = L(4 + k)
    wd = f"WEEKDAY($B$7+{k},1)"
    C[f"{col}1"] = (f'=IF(INDEX(Settings!$C$69:$C$75,{wd})="",INDEX(Settings!$B$69:$B$75,{wd}),'
                    f"INDEX(Settings!$C$69:$C$75,{wd}))")
# 6x7 matrix D3:J8
for w in range(6):
    for d in range(7):
        c = C.cell(3 + w, 4 + d)
        c.value = f"=$B$7+7*{w}+{d}"
        c.number_format = "d mmm"
# month display name (translation-aware) and stamp
C["A10"] = "wk"; C["B9"] = "start"; C["C9"] = "clip start"; C["D9"] = "clip end"; C["E9"] = "active"; C["F9"] = "label"
C["L2"] = "Month name"
C["M2"] = '=IF(INDEX(Settings!$C$79:$C$90,B1)="",INDEX(Settings!$B$79:$B$90,B1),INDEX(Settings!$C$79:$C$90,B1))'
C["L3"] = "Title"; C["M3"] = '=M2&" "&Settings!C4'
for n in range(1, 7):
    r = 9 + n
    C[f"A{r}"] = n
    C[f"B{r}"] = f"=$B$7+7*{n-1}"
    C[f"C{r}"] = f"=MAX(B{r},$B$2)"
    C[f"D{r}"] = f"=MIN(B{r}+6,$B$3)"
    C[f"E{r}"] = f"={n}<=$B$8"
    C[f"F{r}"] = f'=IF(E{r},TEXT(C{r},"mmm d")&" - "&TEXT(D{r},"mmm d"),"Not used this month")'
    for col in "BCD":
        C[f"{col}{r}"].number_format = "d mmm"
C.column_dimensions["F"].width = 22

# =====================================================================
# helpers shared by Month / W tabs
# =====================================================================
def date_fmt(c, fmt="d"): c.number_format = fmt

def weekday_label(date_ref):
    wd = f"WEEKDAY({date_ref},1)"
    return (f'=IF(INDEX(Settings!$C$69:$C$75,{wd})="",INDEX(Settings!$B$69:$B$75,{wd}),'
            f"INDEX(Settings!$C$69:$C$75,{wd}))")

def rept_bar(pct, n=10):
    return f'=IF({pct}="","",REPT("█",ROUND({pct}*{n},0))&REPT("░",{n}-ROUND({pct}*{n},0)))'

def tick_dv(ws):  # helper: nothing – booleans are plain TRUE/FALSE until the setup script converts them
    pass

# =====================================================================
# Month
# =====================================================================
M = ws_mon
widths = {"A": 2, "B": 6, "C": 14, "D": 36, "E": 9, "F": 15, "G": 15, "H": 15, "I": 15, "J": 15, "K": 15, "L": 15, "M": 2, "N": 6}
for k, v in widths.items(): M.column_dimensions[k].width = v
M["B2"] = "=Calc!M3"; M["B2"].font = f(20, True, PRIMARY)
M["F2"] = "Type events directly in the cells under each date. Merge, colour or add emoji as you like."
M["F2"].font = f(9, color=MUTED, italic=True)

banner(M, 3, 2, 4, lab(L_GOALS))
for i in range(5):
    r = 4 + i
    inp(M.cell(r, 2, False), align=Alignment(horizontal="center"))
    inp(M.cell(r, 3)); inp(M.cell(r, 4))
banner(M, 10, 2, 4, lab(L_TODO))
M["D10"] = '=IFERROR(SUMPRODUCT((D11:D25<>"")*(B11:B25=TRUE))/SUMPRODUCT((D11:D25<>"")*1),0)'
M["D10"].number_format = "0%"; M["D10"].alignment = Alignment(horizontal="right")
for i in range(15):
    r = 11 + i
    inp(M.cell(r, 2, False), align=Alignment(horizontal="center"))
    inp(M.cell(r, 3)); inp(M.cell(r, 4))
banner(M, 27, 2, 4, lab(L_DATES))
for i in range(5):
    r = 28 + i
    style(M.cell(r, 2), bg=SURFACE, border=BOX)
    inp(M.cell(r, 3), fmt="d mmm"); inp(M.cell(r, 4))
banner(M, 34, 2, 4, lab(L_NOTES))
for r in range(35, 41):
    for col in (2, 3, 4):
        inp(M.cell(r, col))

# calendar
for k in range(7):
    c = M.cell(4, 6 + k, f"=Calc!{L(4+k)}1")
    style(c, font=f(10, True, "FFFFFF"), bg=PRIMARY, align=Alignment(horizontal="center"))
style(M["E4"], bg=PRIMARY)
for w in range(6):
    top = 5 + 4 * w
    M.cell(top, 5).value = (f'=HYPERLINK("#\'W{w+1}\'!A1","Week {w+1} ▶")')
    style(M.cell(top, 5), font=Font(name=FONT, size=9, color=PRIMARY, underline="single"))
    M.cell(top, 14).value = f"=INDEX(Calc!$E$10:$E$15,{w+1})"
    M.cell(top, 14).font = f(8, color="FFFFFF")
    for d in range(7):
        c = M.cell(top, 6 + d)
        c.value = f'=IF(MONTH(Calc!{L(4+d)}{3+w})=Calc!$B$1,Calc!{L(4+d)}{3+w},"")'
        style(c, font=f(11, True, PRIMARY), bg=MIST, border=BOX, fmt="d", align=Alignment(horizontal="left", vertical="top"))
        for e in range(1, 4):
            inp(M.cell(top + e, 6 + d), align=Alignment(wrap_text=True, vertical="top"))
    for e in range(0, 4):
        M.row_dimensions[top + e].height = 20 if e == 0 else 26
M.freeze_panes = "A3"

# Month CF
today_rng = " ".join(f"F{5+4*w}:L{5+4*w}" for w in range(6))
M.conditional_formatting.add(today_rng, FormulaRule(formula=["AND(ISNUMBER(F5),F5=TODAY())"],
                             fill=fill(ACCENT), font=Font(bold=True, color=TEXT)))
for w in range(6):
    top = 5 + 4 * w
    M.conditional_formatting.add(f"E{top}:L{top+3}", FormulaRule(formula=[f"$N${top}=FALSE"], fill=fill("ECEFF1"),
                                 font=Font(color=DIM)))
M.conditional_formatting.add("D11:D25", FormulaRule(formula=["B11=TRUE"], font=Font(strike=True, color=MUTED)))
M.conditional_formatting.add("C4:D8", FormulaRule(formula=["$B4=TRUE"], font=Font(strike=True, color=MUTED)))
dv = DataValidation(type="list", formula1="=Settings!$B$12:$B$18", allow_blank=True)
M.add_data_validation(dv)
M.column_dimensions["N"].hidden = True
dv.add("C11:C25")   # priority dropdown on the monthly to-do

# =====================================================================
# Weekly tabs
# =====================================================================
BLOCK = 6
def dcol(d, off=0): return 10 + BLOCK * (d - 1) + off          # d = 1..7, off = 0..5
HAB0 = 53                                                       # BA
NUM_HELP0 = 65                                                  # BM

for wi, W in enumerate(ws_w, 1):
    W.column_dimensions["A"].width = 2
    for c in range(2, 9): W.column_dimensions[L(c)].width = 7.5
    W.column_dimensions["I"].width = 3
    for d in range(1, 8):
        W.column_dimensions[L(dcol(d, 0))].width = 11     # time / checkbox
        for off, wdt in ((1, 11), (2, 13), (3, 20), (4, 13), (5, 3)):
            W.column_dimensions[L(dcol(d, off))].width = wdt
    W.column_dimensions[L(HAB0 - 1)].width = 3
    W.column_dimensions[L(HAB0)].width = 26
    for c in range(HAB0 + 1, HAB0 + 8): W.column_dimensions[L(c)].width = 7
    W.column_dimensions[L(HAB0 + 8)].width = 9
    W.column_dimensions[L(HAB0 + 9)].width = 14
    for c in range(NUM_HELP0, NUM_HELP0 + 8): W.column_dimensions[L(c)].hidden = True

    # --- header
    W["B1"] = f'=Calc!M3&"  ·  "&INDEX(Calc!$F$10:$F$15,$B$2)'
    W["B1"].font = f(16, True, PRIMARY)
    W["B2"] = wi; W["B2"].number_format = '"Week "0'; W["B2"].font = f(10, True, MUTED)
    W["H1"] = "=Calc!$B$1"; W["H1"].number_format = ";;;"; W["H1"].font = f(8, color="FFFFFF")
    W["D2"] = '=HYPERLINK("#\'Month\'!A1","◀ Month")'
    W["D2"].font = Font(name=FONT, size=9, color=PRIMARY, underline="single")
    if wi > 1:
        W["F2"] = f'=HYPERLINK("#\'W{wi-1}\'!A1","‹ Prev")'
        W["F2"].font = Font(name=FONT, size=9, color=PRIMARY, underline="single")
    if wi < 6:
        W["H2"] = f'=HYPERLINK("#\'W{wi+1}\'!A1","Next ›")'
        W["H2"].font = Font(name=FONT, size=9, color=PRIMARY, underline="single")
    W["B3"] = '=IF(INDEX(Calc!$E$10:$E$15,$B$2),"","This week is not used in "&Calc!$M$2&".")'
    W["B3"].font = f(10, True, DANGER)

    # --- mini calendar B5:H11
    for k in range(7):
        c = W.cell(5, 2 + k, f"=LEFT(Calc!{L(4+k)}1,2)")
        style(c, font=f(9, True, "FFFFFF"), bg=PRIMARY, align=Alignment(horizontal="center"))
    for w in range(6):
        for k in range(7):
            c = W.cell(6 + w, 2 + k, f"=Calc!{L(4+k)}{3+w}")
            style(c, font=f(9), bg="FFFFFF", border=BOX, fmt="d", align=Alignment(horizontal="center"))
    W.conditional_formatting.add("B6:H11", FormulaRule(formula=["B6=TODAY()"], fill=fill(ACCENT), font=Font(bold=True)))
    W.conditional_formatting.add("B6:H11", FormulaRule(formula=["ROW()-ROW($B$6)+1=$B$2"], fill=fill(MIST),
                                 font=Font(bold=True, color=PRIMARY)))
    W.conditional_formatting.add("B6:H11", FormulaRule(formula=["MONTH(B6)<>$H$1"], font=Font(color=DIM)))

    # --- weekly priorities (B14:H19) / secondary (B22:H27)
    for top, key in ((14, L_WPRI), (22, L_WSEC)):
        banner(W, top, 2, 8, lab(key))
        rng_txt = f"C{top+1}:C{top+5}"
        W.cell(top, 8).value = f'=IFERROR(SUMPRODUCT((C{top+1}:C{top+5}<>"")*(B{top+1}:B{top+5}=TRUE))/SUMPRODUCT((C{top+1}:C{top+5}<>"")*1),0)'
        W.cell(top, 8).number_format = "0%"; W.cell(top, 8).alignment = Alignment(horizontal="right")
        for i in range(5):
            r = top + 1 + i
            inp(W.cell(r, 2, False), align=Alignment(horizontal="center"))
            for col in range(3, 9): inp(W.cell(r, col))
        W.conditional_formatting.add(f"C{top+1}:H{top+5}", FormulaRule(formula=[f"$B{top+1}=TRUE"],
                                     font=Font(strike=True, color=MUTED)))

    # --- day blocks
    for d in range(1, 8):
        c0, c1, c2, c3, c4, c5 = (dcol(d, o) for o in range(6))
        l0, l1, l2, l3, l4, l5 = (L(x) for x in (c0, c1, c2, c3, c4, c5))
        # header
        for col in range(c0, c5):
            style(W.cell(5, col), bg=PRIMARY, font=f(11, True, "FFFFFF"))
        W.cell(5, c0).value = weekday_label(f"${l4}$5")
        W.cell(5, c4).value = f"=INDEX(Calc!$D$3:$J$8,$B$2,{d})"
        W.cell(5, c4).number_format = "mmm d"; W.cell(5, c4).alignment = Alignment(horizontal="right")
        # progress
        task_rng = f"{l3}47:{l3}56"; chk_rng = f"{l0}47:{l0}56"
        calc(W.cell(6, c0, f'=IFERROR(SUMPRODUCT(({task_rng}<>"")*({chk_rng}=TRUE))/SUMPRODUCT(({task_rng}<>"")*1),0)'),
             fmt="0%", align=Alignment(horizontal="center"))
        for col in range(c1, c5): calc(W.cell(6, col))
        W.cell(6, c1).value = rept_bar(f"{l0}6", 14)
        W.cell(6, c1).font = f(10, color=GREEN)
        # schedule
        banner(W, 8, c0, c4, lab(L_SCHED))
        for n in range(36):
            r = 9 + n
            t = W.cell(r, c0)
            t.value = (f'=IF({n}*Settings!$C$7>=1440,"",'
                       f"MOD(ROUND((Settings!$C$6+{n}*Settings!$C$7/1440)*1440,0),1440)/1440)")
            style(t, font=f(9, True, MUTED), bg=SURFACE, border=BOX, fmt="h:mm AM/PM", align=Alignment(horizontal="right"))
            for col in (c1, c2, c3): inp(W.cell(r, col))
            inp(W.cell(r, c4))
        # to-do
        banner(W, 46, c0, c4, lab(L_TODO))
        for i in range(10):
            r = 47 + i
            inp(W.cell(r, c0, False), align=Alignment(horizontal="center"))
            inp(W.cell(r, c1)); inp(W.cell(r, c2)); inp(W.cell(r, c3)); inp(W.cell(r, c4))
        # gratitude
        banner(W, 58, c0, c4, lab(L_GRAT))
        for i in range(5):
            r = 59 + i
            for col in range(c0, c5): inp(W.cell(r, col))
        banner(W, 65, c0, c4, lab(L_NOTES))
        for r in range(66, 72):
            for col in range(c0, c5): inp(W.cell(r, col))
        # CF – header today / out of month
        W.conditional_formatting.add(f"{l0}5:{l4}5", FormulaRule(formula=[f"${l4}$5=TODAY()"], fill=fill(ACCENT),
                                     font=Font(bold=True, color=TEXT)))
        W.conditional_formatting.add(f"{l0}5:{l4}5", FormulaRule(formula=[f"MONTH(${l4}$5)<>$H$1"], fill=fill("7C9AA8")))
        # strike
        W.conditional_formatting.add(f"{l3}47:{l4}56", FormulaRule(formula=[f"${l0}47=TRUE"],
                                     font=Font(strike=True, color=MUTED)))

    # --- dropdowns on day blocks
    dvp = DataValidation(type="list", formula1="=Settings!$B$12:$B$18", allow_blank=True)
    dvc = DataValidation(type="list", formula1="=Settings!$B$22:$B$33", allow_blank=True)
    W.add_data_validation(dvp); W.add_data_validation(dvc)
    for d in range(1, 8):
        dvp.add(f"{L(dcol(d,1))}47:{L(dcol(d,1))}56")
        dvc.add(f"{L(dcol(d,2))}47:{L(dcol(d,2))}56")
        dvc.add(f"{L(dcol(d,4))}9:{L(dcol(d,4))}44")

    # --- slot colours (priorities / categories); relative refs resolve per area
    pr_rng = " ".join(f"{L(dcol(d,1))}47:{L(dcol(d,1))}56" for d in range(1, 8))
    td_cat = " ".join(f"{L(dcol(d,2))}47:{L(dcol(d,2))}56" for d in range(1, 8))
    sc_cat = " ".join(f"{L(dcol(d,4))}9:{L(dcol(d,4))}44" for d in range(1, 8))
    p0, t0, s0 = f"{L(dcol(1,1))}47", f"{L(dcol(1,2))}47", f"{L(dcol(1,4))}9"
    for i in range(7):
        W.conditional_formatting.add(pr_rng, FormulaRule(
            formula=[f'AND({p0}<>"",{p0}=INDIRECT("Settings!$B${12+i}"))'], fill=fill(PRIO_COL[i])))
    for i in range(12):
        for rng, a in ((td_cat, t0), (sc_cat, s0)):
            W.conditional_formatting.add(rng, FormulaRule(
                formula=[f'AND({a}<>"",{a}=INDIRECT("Settings!$B${22+i}"))'], fill=fill(CAT_COL[i])))
    # orphan guard (value no longer in list)
    for rng, a, lst in ((pr_rng, p0, "$B$12:$B$18"), (td_cat, t0, "$B$22:$B$33"), (sc_cat, s0, "$B$22:$B$33")):
        W.conditional_formatting.add(rng, FormulaRule(
            formula=[f'AND({a}<>"",COUNTIF(INDIRECT("Settings!{lst.replace("$","")}"),{a})=0)'],
            font=Font(color=DANGER, bold=True)))

    # --- unused week: grey
    W.conditional_formatting.add("B4:BJ71", FormulaRule(formula=['$B$3<>""'], font=Font(color=DIM)))

    # --- habit tracker
    hb = HAB0
    banner(W, 4, hb, hb + 9, lab(L_HAB))
    style(W.cell(5, hb), bg=MIST, font=f(9, True, PRIMARY))
    for d in range(1, 8):
        c = W.cell(5, hb + d, f"=LEFT({weekday_label('INDEX(Calc!$D$3:$J$8,$B$2,%d)' % d)[1:]},3)")
        style(c, bg=MIST, font=f(9, True, PRIMARY), align=Alignment(horizontal="center"))
    style(W.cell(5, hb + 8, "%"), bg=MIST, font=f(9, True, PRIMARY), align=Alignment(horizontal="center"))
    style(W.cell(5, hb + 9, "Progress"), bg=MIST, font=f(9, True, PRIMARY))
    a, bcol, hcol, pcol, barcol = L(hb), L(hb + 1), L(hb + 7), L(hb + 8), L(hb + 9)
    for i in range(21):
        r = 6 + i
        calc(W.cell(r, hb, f'=IF(Settings!$B${37+i}="","",Settings!$B${37+i})'))
        for d in range(1, 8):
            inp(W.cell(r, hb + d, False), align=Alignment(horizontal="center"))
        calc(W.cell(r, hb + 8, f'=IF(${a}{r}="","",COUNTIF({bcol}{r}:{hcol}{r},TRUE)/7)'),
             fmt="0%", align=Alignment(horizontal="center"))
        calc(W.cell(r, hb + 9, rept_bar(f"{pcol}{r}", 10)))
        W.cell(r, hb + 9).font = f(9, color=GREEN)
    calc(W.cell(27, hb, "Daily completion"))
    W.cell(27, hb).font = f(9, True, PRIMARY)
    for d in range(1, 8):
        col = L(hb + d)
        calc(W.cell(27, hb + d, f'=IFERROR(SUMPRODUCT((${a}$6:${a}$26<>"")*({col}$6:{col}$26=TRUE))/SUMPRODUCT((${a}$6:${a}$26<>"")*1),0)'),
             fmt="0%", align=Alignment(horizontal="center"))
    calc(W.cell(27, hb + 8, f'=IFERROR(SUMPRODUCT((${a}$6:${a}$26<>"")*(${bcol}$6:${hcol}$26=TRUE))/(SUMPRODUCT((${a}$6:${a}$26<>"")*1)*7),0)'),
         fmt="0%", align=Alignment(horizontal="center"))
    calc(W.cell(27, hb + 9, rept_bar(f"{pcol}27", 10))); W.cell(27, hb + 9).font = f(9, color=GREEN)
    W.conditional_formatting.add(f"{bcol}6:{hcol}26", FormulaRule(formula=[f"{bcol}6=TRUE"], fill=fill("CDEBDD")))
    # hide habit rows with no label (grey out)
    W.conditional_formatting.add(f"{a}6:{barcol}26", FormulaRule(formula=[f'${a}6=""'], fill=fill("F2F4F5"), font=Font(color=DIM)))

    # --- numeric trackers (rows 30..)
    banner(W, 30, hb, hb + 9, lab(L_NUM))
    style(W.cell(30, hb + 8), bg=PRIMARY, font=f(9, True, "FFFFFF")); W.cell(30, hb + 8).value = "Avg"
    style(W.cell(30, hb + 9), bg=PRIMARY, font=f(9, True, "FFFFFF")); W.cell(30, hb + 9).value = "Total"
    for t in range(5):
        r = 31 + 2 * t; sr = 61 + t
        calc(W.cell(r, hb, f'=Settings!$B${sr}&IF(Settings!$C${sr}<>""," ("&Settings!$C${sr}&")","")'))
        W.cell(r, hb).font = f(10, True)
        calc(W.cell(r + 1, hb, "progress"))
        W.cell(r + 1, hb).font = f(8, color=MUTED, italic=True)
        for d in range(1, 8):
            col = L(hb + d); hcol_ = L(NUM_HELP0 + d - 1)
            inp(W.cell(r, hb + d), align=Alignment(horizontal="center"))
            # hidden fraction (min/max from Settings)
            W.cell(r, NUM_HELP0 + d - 1).value = (f'=IF({col}{r}="","",IFERROR(MIN(1,MAX(0,({col}{r}-Settings!$D${sr})/'
                                                  f"(Settings!$E${sr}-Settings!$D${sr}))),0))")
            calc(W.cell(r + 1, hb + d, rept_bar(f"{hcol_}{r}", 5)), align=Alignment(horizontal="center"))
            W.cell(r + 1, hb + d).font = f(8, color=GREEN)
        first, last = L(hb + 1), L(hb + 7)
        calc(W.cell(r, hb + 8, f"=IFERROR(AVERAGE({first}{r}:{last}{r}),0)"), fmt="0.0", align=Alignment(horizontal="center"))
        calc(W.cell(r, hb + 9, f"=SUM({first}{r}:{last}{r})"), fmt="0.0", align=Alignment(horizontal="center"))
        W.cell(r, NUM_HELP0 + 7).value = (f"=IFERROR(MIN(1,MAX(0,({L(hb+8)}{r}-Settings!$D${sr})/(Settings!$E${sr}-Settings!$D${sr}))),0)")
        W.cell(r + 1, hb + 8).value = rept_bar(f"{L(NUM_HELP0+7)}{r}", 5)
        style(W.cell(r + 1, hb + 8), font=f(8, color=GREEN), bg=SURFACE, border=BOX, align=Alignment(horizontal="center"))
        # validation: numeric within Min..Max
        dvn = DataValidation(type="decimal", operator="between", formula1=f"Settings!$D${sr}",
                             formula2=f"Settings!$E${sr}", allow_blank=True, showErrorMessage=True,
                             errorTitle="Out of range", error="Value must be between Min and Max set on the Settings tab.")
        W.add_data_validation(dvn); dvn.add(f"{first}{r}:{last}{r}")

    # --- freeze / print / protection
    W.freeze_panes = "A4"
    W.sheet_properties.tabColor = PRIMARY if wi % 2 else "4A7C8F"
    W.protection.sheet = True
    for flag in ("formatCells", "formatColumns", "formatRows", "insertRows", "deleteRows", "sort", "autoFilter"):
        setattr(W.protection, flag, False)

M.protection.sheet = True
for flag in ("formatCells", "formatColumns", "formatRows", "sort"):
    setattr(M.protection, flag, False)
M.sheet_properties.tabColor = ACCENT
S.sheet_properties.tabColor = "6B7A83"
S.protection.sheet = True
for flag in ("formatCells",):
    setattr(S.protection, flag, False)
ws_calc.sheet_state = "hidden"; ws_lists.sheet_state = "hidden"
ws_calc.protection.sheet = True; ws_lists.protection.sheet = True

# =====================================================================
# Start
# =====================================================================
T = ws_start
T.column_dimensions["A"].width = 2; T.column_dimensions["B"].width = 110
T["B2"] = "Monthly · Weekly · Daily Planner"; T["B2"].font = f(22, True, PRIMARY)
T["B3"] = "Original planner built from the blueprint in BLUEPRINT.md. Reuse it every month."; T["B3"].font = f(10, color=MUTED)
steps = [
    ("1  Keep one clean master copy.", "Each month: File > Make a copy, then work in the copy. Typed events are not tied to dates."),
    ("2  Open Settings.", "Pick the month, year, week start, schedule start time and slot length. All dates and times update by themselves."),
    ("3  Fill the lists.", "Priorities (7), categories (12), habits (21) and the five numeric trackers. They appear in every dropdown on every week."),
    ("4  Plan the month.", "On Month: goals, monthly to-do, important dates; type events in the cells under each date."),
    ("5  Plan each week (W1 to W6).", "Weekly priorities, daily schedule, daily to-do, gratitude, notes. Tick a task and its text is struck through and the bar moves."),
    ("6  Track.", "Tick habits and enter numbers on the far right of each week tab."),
    ("", ""),
    ("Good to know", ""),
    ("•", "Grey cells are formulas – don't type over them. White cells are yours. You may merge cells in white areas."),
    ("•", "Days from the neighbouring month appear dimmed in a week tab. Track such a day in only one month's copy."),
    ("•", "Weeks the month doesn't need show a red note and turn grey."),
    ("•", "Google Sheets: run the setup script once (setup_google_sheets.gs) to turn TRUE/FALSE cells into checkboxes and the text bars into coloured bars."),
    ("•", "Changing the month AFTER you typed events leaves your text on the old cells – make a new copy instead."),
]
for i, (a, b) in enumerate(steps):
    r = 5 + i
    T.cell(r, 2, (a + ("  " + b if b else "")))
    T.cell(r, 2).font = f(11, bool(a and not b) or a[:1].isdigit() and False)
    T.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
    if a and a[:1].isdigit():
        T.cell(r, 2).font = f(11)
T["B12"].font = f(12, True, PRIMARY)
T.sheet_properties.tabColor = PRIMARY
T.protection.sheet = True

wb.defined_names  # (names kept as labels in Calc!A; formulas use direct refs for portability)
wb.save(OUT)
print("saved", OUT)
